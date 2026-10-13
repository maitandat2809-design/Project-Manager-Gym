from copy import deepcopy
from pathlib import Path
from docx import Document

SOURCE = Path(r"C:\Users\lamtu\Desktop\Final Edit - Proper Table of Figures.docx")
OUTPUT = Path(r"C:\xampp\htdocs\Gym Manager\Final Edit - Objectives Updated.docx")


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def insert_after(paragraph, text, style, template):
    new_paragraph = paragraph.insert_paragraph_before()
    paragraph._p.addnext(new_paragraph._p)
    new_paragraph.style = style
    new_paragraph.paragraph_format._element.getparent().remove(new_paragraph.paragraph_format._element) if False else None
    run = new_paragraph.add_run(text)
    if template.runs:
        run.font.name = template.runs[0].font.name
        run.font.size = template.runs[0].font.size
    return new_paragraph


document = Document(SOURCE)
paragraphs = document.paragraphs
anchor = paragraphs[145]
normal_template = paragraphs[146]
heading3_template = paragraphs[147]

# First remove the old condensed objective material. This keeps the chapter heading
# and the following 1.4 methodology section unchanged.
for paragraph in list(document.paragraphs[146:163]):
    remove_paragraph(paragraph)

# Each tuple is (heading level, text).  The outline exactly follows the four-objective
# structure of the approved proposal, while the prose reflects the implemented system.
content = [
    (0, "This section restates the four objectives in the same structure as the approved proposal and updates them to reflect the completed Gym Manager Website, its four user roles, and the PHP/MySQL implementation."),
    (3, "1.3.1 Develop a Gym Management and Online Appointment Scheduling System"),
    (4, "1.3.1.1 Detailed Description of the Objective"),
    (0, "The primary objective is to deliver one integrated Gym Manager Website for customers, coaches, gym managers, and administrators. The system combines secure account access, coach discovery, one-to-one training booking, membership management, online product purchasing, payment handling, equipment information, and administrative operations on one platform. The objective is achieved when users can complete registration, login, booking, and purchasing workflows reliably, while authorised staff can manage operational data without inconsistency."),
    (4, "1.3.1.2 Analysis and Design of System Requirements"),
    (0, "The requirements and design stage defines the four system roles, their permissions, use cases, activity flows, site maps, class diagrams, and the MySQL entity-relationship design. It covers the data needed for users, coaches, bookings, memberships, products, orders, payments, reviews, password resets, and equipment. This stage is successful when the documented requirements are traceable to the system design and the database structure supports the booking, subscription, shopping, and management workflows."),
    (0, "Estimated time: Weeks 1-8."),
    (4, "1.3.1.3 Development and Implementation of Core Features"),
    (0, "Implementation delivers account registration, login, password recovery by email, personal profile management, coach information, conflict-checked 1-to-1 booking, booking confirmation or cancellation, attendance recording, membership selection, online shopping, cart and checkout, order tracking, product and coach reviews, gym equipment information, and role-based management functions. Payment options include MoMo and cash on delivery (COD), and relevant authentication, booking, cancellation, and payment events trigger SMTP email notifications."),
    (0, "The objective is successful when the planned core workflows operate without critical errors, registration and login tests achieve at least a 95% pass rate, and functional testing achieves at least a 90% pass rate for the documented requirements."),
    (0, "Estimated time: Weeks 9-16."),
    (4, "1.3.1.4 Testing and System Improvement"),
    (0, "The completed system is tested through functional, integration, role-permission, validation, and usability tests. Testing verifies that invalid inputs, booking conflicts, unauthorised access, inconsistent payments, and stock-related errors are prevented or reported clearly. Improvements focus on fixing critical defects, strengthening validation, improving messages, and making the booking and checkout journeys easier to use."),
    (0, "The objective is achieved when critical defects are resolved, normal requests respond in approximately three seconds or less under the project test conditions, and at least 85% of test users can complete key tasks without assistance."),
    (3, "1.3.2 System Users and Requirements of the Gym Management Platform"),
    (4, "1.3.2.1 Identification of System Users and Objectives"),
    (0, "The platform has four clearly separated roles: User/Customer, Coach, Gym Manager, and Administrator. Users maintain their profiles, book sessions, manage memberships, purchase products, follow orders, view equipment, and use the AI assistant. Coaches manage their professional information, schedules, attendance, and training notes. Gym Managers manage subscriptions, products, orders, bookings, equipment, notifications, and revenue reports. Administrators manage user, coach, and manager accounts, including roles and permissions."),
    (0, "This objective is achieved when every role has a defined responsibility and can access only the functions appropriate to that role."),
    (4, "1.3.2.2 Analysis of User Requirements"),
    (0, "User requirements are analysed as role-specific functional requirements. Customers require simple registration, profile maintenance, coach and schedule viewing, booking, membership purchase, shopping, order tracking, and notification support. Coaches require a calendar, session approval, attendance tools, and notes. Gym Managers require operational control of memberships, inventory, orders, bookings, equipment, and reports. Administrators require account, coach, manager, and permission administration."),
    (0, "The objective is achieved when each prioritised user need is documented and mapped to a specific feature, interface, and database record."),
    (4, "1.3.2.3 Development of Features Based on User Needs"),
    (0, "The system implements these requirements through role-based access control, server-side sessions, password hashing, input validation, prepared database statements, and separate role workspaces. Core workflows have been verified using the relevant role account, including customer booking and checkout, coach attendance and notes, manager operations, and administrator account management."),
    (0, "This objective is achieved when authorised users can complete their main duties successfully and unauthorised cross-role access is rejected."),
    (3, "1.3.3 Improving User Satisfaction and Management Efficiency"),
    (4, "1.3.3.1 Objective Description"),
    (0, "This objective evaluates how the implemented platform improves the customer experience and administrative efficiency compared with fragmented manual processes. Customers should obtain information, book services, buy products, and review histories more conveniently. Staff should manage connected booking, membership, order, inventory, equipment, and reporting data from one system instead of duplicate records."),
    (0, "The objective is achieved if the platform reduces avoidable data-entry errors, shortens routine task completion, and gives users a clear and understandable experience."),
    (4, "1.3.3.2 Enhancing User Satisfaction"),
    (0, "The website uses responsive pages for desktop, tablet, and mobile browsers, bilingual navigation, clear labels, validation messages, and streamlined booking and checkout processes. Customers can view booking, membership, and order histories, while SMTP notifications confirm important events such as password recovery, payment, booking confirmation, and booking cancellation."),
    (0, "Success is measured through usability testing: at least 85% of test users should complete a booking or purchase without assistance, and the average usability rating should be at least 4 out of 5."),
    (4, "1.3.3.3 Improving Management Efficiency"),
    (0, "Centralised MySQL records allow managers to work from shared data for memberships, bookings, products, orders, equipment, and revenue. Validation and status controls prevent invalid booking times, incomplete payment information, and inconsistent order handling. Monthly and yearly reports provide information for routine operational decisions."),
    (0, "The objective is achieved when defined manual-versus-system scenarios show at least an 80% reduction in duplicated or inconsistent data-entry outcomes."),
    (3, "1.3.4 Expanding System Functionality to Support Business Operations"),
    (4, "1.3.4.1 Objective Description"),
    (0, "This objective expands the gym management platform beyond training administration to support membership revenue and online sports-product sales. The same system records customer transactions, order data, payment status, and business information so that staff can operate services and commerce without switching to a separate platform."),
    (0, "The objective is achieved when sales and service functions work together without compromising data consistency or usability."),
    (4, "1.3.4.2 Implementation of Online Sports Equipment Sales"),
    (0, "Customers can browse and filter products, inspect product details, select variants where available, add items to a cart, check out, select MoMo or COD payment, receive payment-related email notifications, and track orders. They can also submit product reviews after use. These functions make gym-related products available through the same account used for memberships and bookings."),
    (0, "The objective is achieved when customers can complete the end-to-end product selection, cart, payment, and order-history flow and each transaction is recorded accurately."),
    (4, "1.3.4.3 Management of Product and Order Data"),
    (0, "Gym Managers can create, edit, and remove product information, maintain stock and weight variants, view orders and invoice details, update order status, cancel or restore eligible orders, and export invoice output. Membership subscriptions are also managed through plan creation, approval or rejection, payment confirmation, and status tracking."),
    (0, "The objective is achieved when product, stock, order-item, subscription, and payment records remain consistent during normal management operations."),
    (4, "1.3.4.4 Monitoring Revenue and Business Performance"),
    (0, "The platform provides Gym Managers with monthly and yearly revenue reports and export options. These reports combine recorded membership and order activity so managers can monitor transaction status, business performance, and operational trends from the same platform."),
    (0, "The objective is achieved when report totals reproduce the values stored in the operational database and managers can access the information needed for routine decisions."),
]

last = anchor
for level, text in content:
    if level == 3:
        last = insert_after(last, text, "Heading 3", heading3_template)
    elif level == 4:
        last = insert_after(last, text, "Heading 4", heading3_template)
    else:
        last = insert_after(last, text, "Normal", normal_template)

document.save(OUTPUT)
print(OUTPUT)
