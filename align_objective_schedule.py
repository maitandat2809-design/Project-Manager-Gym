from pathlib import Path
from docx import Document

SOURCE = Path(r"C:\Users\lamtu\Desktop\Final Edit - Objectives Updated.docx")
OUTPUT = Path(r"C:\xampp\htdocs\Gym Manager\Final Edit - Objectives 24-Week Schedule Updated.docx")


def insert_after(paragraph, text):
    new_paragraph = paragraph.insert_paragraph_before()
    paragraph._p.addnext(new_paragraph._p)
    new_paragraph.style = "Normal"
    new_paragraph.add_run(text)
    return new_paragraph


document = Document(SOURCE)

# The overall statement makes it clear that Appendix A, not the proposal estimates,
# is the authoritative 24-week project schedule.
for paragraph in document.paragraphs:
    if paragraph.text.startswith("This section restates the four objectives"):
        paragraph.text = (
            "This section restates the four objectives in the same structure as the approved proposal "
            "and updates them to reflect the completed Gym Manager Website, its four user roles, and "
            "the PHP/MySQL implementation. All time allocations below follow the 24-week Waterfall "
            "schedule presented in Appendix A."
        )
        paragraph.style = "Normal"
        break

replacements = {
    "Estimated time: Weeks 1-8.": "Schedule allocation: Weeks 1-8 (Requirements Analysis and System Design).",
    "Estimated time: Weeks 9-16.": "Schedule allocation: Weeks 9-16 (Implementation).",
}
for paragraph in document.paragraphs:
    if paragraph.text in replacements:
        paragraph.text = replacements[paragraph.text]
        paragraph.style = "Normal"

# Give every main objective an explicit Appendix A schedule allocation. The wording
# retains the detailed phases where they occur and avoids using the old proposal estimates.
schedule_after = {
    "The objective is achieved when critical defects are resolved, normal requests respond in approximately three seconds or less under the project test conditions, and at least 85% of test users can complete key tasks without assistance.":
        "Schedule allocation: Weeks 17-20 (Testing and defect correction). Overall, Objective 1 spans Weeks 1-20.",
    "This objective is achieved when authorised users can complete their main duties successfully and unauthorised cross-role access is rejected.":
        "Schedule allocation: Weeks 1-8 for role analysis and design, Weeks 10-11 and 16 for role-specific implementation, and Weeks 17-18 for role and access-permission testing.",
    "The objective is achieved when defined manual-versus-system scenarios show at least an 80% reduction in duplicated or inconsistent data-entry outcomes.":
        "Schedule allocation: Weeks 9-20 for the usability-related implementation and testing work, followed by Week 23 for usability evaluation and analysis.",
    "The objective is achieved when report totals reproduce the values stored in the operational database and managers can access the information needed for routine decisions.":
        "Schedule allocation: Weeks 13-16 for membership, shop, payment, and management implementation; Weeks 19-20 for integration and correction; and Weeks 21-22 for deployment and final operational checks.",
}

for paragraph in list(document.paragraphs):
    if paragraph.text in schedule_after:
        insert_after(paragraph, schedule_after[paragraph.text])

document.save(OUTPUT)
print(OUTPUT)
